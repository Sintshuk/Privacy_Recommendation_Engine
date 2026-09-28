"""
Turns raw extraction output into a clean dataset: csv, json and exel sheet. 
"""

from __future__ import annotations

import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

COLUMNS = ["policy", "segment_id", "id", "attribute", "value", "evidence", "confidence", "linked_ids", "link"]
CSV_COLUMNS = ["policy", "segment_id", "id", "attribute", "value", "evidence", "segment_text", "confidence", "linked_ids", "link"]
SEGMENT_COLUMNS = ["policy", "segment_id", "segment_text"]

HEADER_FILL_COLOR = "1F2937"
HEADER_FONT_COLOR = "FFFFFF"


def flatten_item(item: dict) -> list[dict]:
    policy = item.get("policy", "")
    segment_id = item.get("segment", {}).get("id", "")
    result = item.get("result") or {}

    rows: list[dict] = []

    for attribute in ("right_type", "choice_type"):
        for entry in result.get(attribute, []):
            rows.append({
                "policy": policy,
                "segment_id": segment_id,
                "id": entry.get("id", ""),
                "attribute": attribute,
                "value": entry.get("value", ""),
                "evidence": entry.get("evidence", ""),
                "confidence": entry.get("confidence", "extracted"),
            })

    for entry in result.get("request_channel", []):
        row = {
            "policy": policy,
            "segment_id": segment_id,
            "id": entry.get("id", ""),
            "attribute": "request_channel",
            "value": entry.get("value", ""),
            "evidence": entry.get("evidence", ""),
            "confidence": entry.get("confidence", "extracted"),
            "linked_ids": entry.get("linked_ids", []),
        }
        link = entry.get("link")
        if link and link.get("url"):
            row["link"] = link["url"]
        rows.append(row)

    return rows


def collect_segment_text(item: dict) -> dict | None:
    segment = item.get("segment", {})
    segment_id = segment.get("id", "")
    text = segment.get("text", "")
    if not segment_id or not text:
        return None
    return {"policy": item.get("policy", ""), "segment_id": segment_id, "segment_text": text}


def build_clean_dataset(raw_data: dict) -> tuple[list[dict], list[dict], list[dict]]:
    """Returns (extractions, segments, skipped). segments is deduped by (policy, segment_id)."""
    extractions: list[dict] = []
    skipped: list[dict] = []
    segments: dict[tuple[str, str], dict] = {}

    for item in raw_data.get("items", []):
        rows = flatten_item(item)
        if rows:
            extractions.extend(rows)
        else:
            result = item.get("result") or {}
            skipped.append({
                "policy": item.get("policy", ""),
                "segment_id": item.get("segment", {}).get("id", ""),
                "reference_count": len(result.get("references", [])),
            })

        segment_row = collect_segment_text(item)
        if segment_row:
            key = (segment_row["policy"], segment_row["segment_id"])
            segments.setdefault(key, segment_row)

    return extractions, list(segments.values()), skipped


def build_grouped_json(extractions: list[dict], segments: list[dict]) -> list[dict]:
    """
    One block per segment: segment_text once, followed by every
    extraction that came from it. Reads top to bottom in document
    order, and segment_text still appears exactly once no matter how
    many extractions a segment produced.
    """
    extractions_by_segment: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for row in extractions:
        key = (row["policy"], row["segment_id"])
        # policy/segment_id already live on the parent block - repeating
        # them on every extraction inside it would be the same
        # redundancy this structure exists to avoid.
        entry = {k: v for k, v in row.items() if k not in ("policy", "segment_id")}
        extractions_by_segment[key].append(entry)

    groups: list[dict] = []
    for segment in segments:
        key = (segment["policy"], segment["segment_id"])
        rows = extractions_by_segment.get(key, [])
        if not rows:
            continue
        groups.append({
            "policy": segment["policy"],
            "segment_id": segment["segment_id"],
            "segment_text": segment["segment_text"],
            "extractions": rows,
        })

    return groups


def write_json(groups: list[dict], path: Path) -> None:
    path.write_text(json.dumps(groups, ensure_ascii=False, indent=2), encoding="utf-8")


def write_csv_rows(rows: list[dict], columns: list[str], path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            flat = dict(row)
            if "linked_ids" in flat:
                flat["linked_ids"] = ";".join(flat["linked_ids"]) if flat["linked_ids"] else ""
            writer.writerow(flat)


def write_csv(extractions: list[dict], segments: list[dict], path: Path) -> None:
    text_lookup = {(s["policy"], s["segment_id"]): s["segment_text"] for s in segments}
    joined = []
    for row in extractions:
        flat = dict(row)
        flat["segment_text"] = text_lookup.get((row.get("policy", ""), row.get("segment_id", "")), "")
        joined.append(flat)
    write_csv_rows(joined, CSV_COLUMNS, path)


def safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r'[:\\/?*\[\]]', "_", name)[:31] or "Policy"
    candidate = cleaned
    suffix = 2
    while candidate in used:
        candidate = f"{cleaned[:31 - len(f'_{suffix}')]}_{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def style_header(sheet, font, fill, alignment) -> None:
    for cell in sheet[1]:
        cell.font = font
        cell.fill = fill
        cell.alignment = alignment
    sheet.freeze_panes = "A2"


def write_excel(extractions: list[dict], segments: list[dict], path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  (skipping Excel export - run: pip install openpyxl)")
        return

    by_policy: dict[str, list[dict]] = {}
    for row in extractions:
        by_policy.setdefault(row.get("policy", "Unknown"), []).append(row)

    segment_texts: dict[tuple[str, str], str] = {
        (s["policy"], s["segment_id"]): s["segment_text"] for s in segments
    }

    workbook = Workbook()
    workbook.remove(workbook.active)

    font = Font(bold=True, color=HEADER_FONT_COLOR)
    fill = PatternFill(start_color=HEADER_FILL_COLOR, end_color=HEADER_FILL_COLOR, fill_type="solid")
    header_alignment = Alignment(vertical="center")
    wrap_alignment = Alignment(vertical="top", wrap_text=True)

    used_names: set[str] = set()

    segments_sheet = workbook.create_sheet(title=safe_sheet_name("Segments", used_names))
    segments_sheet.append(SEGMENT_COLUMNS)
    style_header(segments_sheet, font, fill, header_alignment)

    for (policy, segment_id), text in sorted(segment_texts.items()):
        segments_sheet.append([policy, segment_id, text])

    segment_widths = {"policy": 14, "segment_id": 11, "segment_text": 100}
    for index, column_name in enumerate(SEGMENT_COLUMNS, start=1):
        segments_sheet.column_dimensions[get_column_letter(index)].width = segment_widths.get(column_name, 20)
    for row_cells in segments_sheet.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = wrap_alignment
        segments_sheet.row_dimensions[row_cells[0].row].height = 90

    for policy_name, policy_rows in by_policy.items():
        sheet = workbook.create_sheet(title=safe_sheet_name(policy_name, used_names))
        sheet.append(COLUMNS)
        style_header(sheet, font, fill, header_alignment)

        for row in policy_rows:
            linked_ids = row.get("linked_ids", [])
            sheet.append([
                row.get("policy", ""), row.get("segment_id", ""), row.get("id", ""),
                row.get("attribute", ""), row.get("value", ""), row.get("evidence", ""),
                row.get("confidence", ""),
                ";".join(linked_ids) if linked_ids else "",
                row.get("link", ""),
            ])

        widths = {"policy": 14, "segment_id": 11, "id": 22, "attribute": 15,
                  "value": 30, "evidence": 70, "confidence": 11, "linked_ids": 30, "link": 40}
        for index, column_name in enumerate(COLUMNS, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = widths.get(column_name, 20)
        for row_cells in sheet.iter_rows(min_row=2):
            for cell in row_cells:
                cell.alignment = wrap_alignment
            sheet.row_dimensions[row_cells[0].row].height = 60

    if len(workbook.sheetnames) == 1:
        workbook.create_sheet(title="extractions")

    workbook.move_sheet("Segments", offset=-(len(workbook.sheetnames) - 1))
    workbook.save(path)


def export_all(raw_data: dict, output_prefix: Path) -> dict:
    """
    The importable entry point - main.py calls this directly after
    every policy so the clean files are never more than one policy
    stale. Returns a small summary dict for logging.
    """
    extractions, segments, skipped = build_clean_dataset(raw_data)
    grouped = build_grouped_json(extractions, segments)

    write_json(grouped, output_prefix.with_name(output_prefix.name + "_clean.json"))
    write_csv(extractions, segments, output_prefix.with_name(output_prefix.name + "_clean.csv"))
    write_excel(extractions, segments, output_prefix.with_name(output_prefix.name + "_clean.xlsx"))

    if skipped:
        skipped_path = output_prefix.with_name(output_prefix.name + "_skipped_segments.json")
        skipped_path.write_text(json.dumps(skipped, ensure_ascii=False, indent=2), encoding="utf-8")

    return {
        "extractions": len(extractions),
        "segments": len(segments),
        "skipped": len(skipped),
        "policies": sorted({row["policy"] for row in extractions}),
    }


def main() -> None:
    if len(sys.argv) < 2:
        print("Usage: python export_clean.py <input_extractions.json> [output_prefix]")
        sys.exit(1)

    input_path = Path(sys.argv[1])
    output_prefix = Path(sys.argv[2]) if len(sys.argv) > 2 else input_path.with_suffix("")

    raw_data = json.loads(input_path.read_text(encoding="utf-8"))
    summary = export_all(raw_data, output_prefix)

    print(f"extractions: {summary['extractions']}  segments: {summary['segments']}  "
          f"skipped: {summary['skipped']}  policies: {len(summary['policies'])}")
    print(f"written: {output_prefix.name}_clean.json / .csv / .xlsx")


if __name__ == "__main__":
    main()